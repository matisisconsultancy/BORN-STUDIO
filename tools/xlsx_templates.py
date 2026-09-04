#!/usr/bin/env python3
"""BORN Studio — redesigned Sourcing & Development trackers (branded XLSX).

Rebuilds the two supplier templates in BORN's visual language: ink/paper/Rojo
palette, a branded banner with the wordmark, grouped headers, row banding,
frozen panes, auto-filter, dropdowns, conditional status formatting, currency
formats, an example row, and a 'Start here' guide sheet.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    RICH=True
except Exception:
    RICH=False
import os

OUTDIR="assets/templates"; os.makedirs(OUTDIR, exist_ok=True)

# ---- brand palette (ARGB) -------------------------------------------------
INK="FF1B1720"; PAPER="FFF2EEE6"; PAPER2="FFEAE5DA"; PAPER3="FFDED8CC"
G2="FF4A454F"; G3="FF8A8590"; G4="FFB7B2BA"; RED="FFC4122E"; WHITE="FFFFFFFF"
INKD="FF14111A"
SERIF="Georgia"; SANS="Arial"; MONO="Courier New"

def fill(c): return PatternFill("solid", fgColor=c)
def side(c, style="thin"): return Side(style=style, color=c)
THIN=side(G4); THINK=side("FFCFC9BE")
def box(l=None,r=None,t=None,b=None): return Border(left=l,right=r,top=t,bottom=b)

def set_widths(ws, widths):  # widths: {colidx: width}
    for i,w in widths.items(): ws.column_dimensions[L(i)].width=w

def merge_center(ws, r1,c1,r2,c2, value, font, fillc=None, align=None):
    ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
    cell=ws.cell(r1,c1); cell.value=value; cell.font=font
    if fillc:
        for r in range(r1,r2+1):
            for c in range(c1,c2+1): ws.cell(r,c).fill=fill(fillc)
    cell.alignment=align or Alignment(horizontal="left", vertical="center")
    return cell

def wordmark(ws, r1,c1,r2,c2, size=22):
    """BORN. — white with a Rojo period (rich text where supported)."""
    ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1): ws.cell(r,c).fill=fill(INK)
    cell=ws.cell(r1,c1)
    if RICH:
        cell.value=CellRichText(
            TextBlock(InlineFont(rFont=SERIF, sz=size, b=True, color="FFFFFF"), "BORN"),
            TextBlock(InlineFont(rFont=SERIF, sz=size, b=True, color="C4122E"), "."))
    else:
        cell.value="BORN."; cell.font=Font(name=SERIF,size=size,bold=True,color=WHITE)
    cell.alignment=Alignment(horizontal="left", vertical="center", indent=1)

# ---------------------------------------------------------------- banner
def banner(ws, ncols, title, kicker):
    ws.sheet_view.showGridLines=False
    wend=min(ncols, 6)
    ws.row_dimensions[1].height=40
    wordmark(ws,1,1,1,wend)
    merge_center(ws,1,wend+1,1,ncols,title,
                 Font(name=MONO,size=12,bold=True,color=WHITE),INK,
                 Alignment(horizontal="right",vertical="center",indent=2))
    # red ruler rule
    ws.row_dimensions[2].height=5
    for c in range(1,ncols+1): ws.cell(2,c).fill=fill(RED)
    # tagline row
    ws.row_dimensions[3].height=18
    merge_center(ws,3,1,3,wend,"From idea to life",
                 Font(name=SERIF,size=10,italic=True,color=G2),PAPER)
    merge_center(ws,3,wend+1,3,ncols,kicker,
                 Font(name=MONO,size=8,color=G3),PAPER,
                 Alignment(horizontal="right",vertical="center",indent=2))
    ws.row_dimensions[4].height=6
    for c in range(1,ncols+1): ws.cell(4,c).fill=fill(PAPER)

# ---------------------------------------------------------------- headers
def group_header(ws, row, groups):
    ws.row_dimensions[row].height=22
    for c1,c2,label in groups:
        merge_center(ws,row,c1,row,c2,label,
                     Font(name=SERIF,size=10,bold=True,color=WHITE),INK,
                     Alignment(horizontal="center",vertical="center"))
        # red accent under each group + ink separators
        for c in range(c1,c2+1):
            ws.cell(row,c).border=box(b=side(RED,"medium"),
                                      l=side(INKD) if c==c1 else None,
                                      r=side(INKD) if c==c2 else None)

def sub_header(ws, row, cols):
    ws.row_dimensions[row].height=40
    for i,(name,*rest) in enumerate(cols, start=1):
        c=ws.cell(row,i); c.value=name
        c.fill=fill(PAPER3)
        c.font=Font(name=SANS,size=9,bold=True,color=INK)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        c.border=box(b=side(G3),l=THINK,r=THINK)

# ---------------------------------------------------------------- data body
def data_body(ws, first_row, nrows, cols, example):
    ncols=len(cols)
    for ridx in range(nrows):
        row=first_row+ridx
        band = PAPER if ridx%2==0 else PAPER2
        ws.row_dimensions[row].height=20
        for i,(name,width,numfmt,kind) in enumerate(cols, start=1):
            c=ws.cell(row,i)
            c.fill=fill(band)
            c.border=box(b=THIN,l=THINK,r=THINK)
            c.font=Font(name=(MONO if kind in("num","cost") else SANS),size=10,color=INK)
            if kind in ("num","cost"):
                c.alignment=Alignment(horizontal="center",vertical="center")
            else:
                c.alignment=Alignment(horizontal="left",vertical="center",indent=1)
            if numfmt: c.number_format=numfmt
            if ridx==0 and (name in example): c.value=example[name]
        # example row accent
        if ridx==0:
            for i in range(1,ncols+1):
                cur=ws.cell(first_row,i).border
                ws.cell(first_row,i).border=box(b=side(G3),l=cur.left,r=cur.right,t=side(G3))
            ws.cell(first_row,1).comment=Comment(
                "Example row — overwrite with your first supplier, or delete it.\nEvery row below is yours to fill.","BORN Studio")

def col_letter_by_name(cols, name):
    for i,(n,*_ ) in enumerate(cols,start=1):
        if n==name: return L(i)
    return None

# ---------------------------------------------------------------- build one tracker
def build_tracker(path, sheet_title, banner_title, kicker, groups, cols, example,
                  category_at=None, eval_at=None, glossary=None, nrows=40):
    wb=openpyxl.Workbook()
    guide=wb.active; guide.title="Start here"
    ws=wb.create_sheet(sheet_title)
    ncols=len(cols)
    HDR_G=5; HDR_S=6; DATA0=7
    banner(ws, ncols, banner_title, kicker)
    group_header(ws, HDR_G, groups)
    sub_header(ws, HDR_S, cols)
    data_body(ws, DATA0, nrows, cols, example)
    set_widths(ws, {i:w for i,(_,w,_,_) in enumerate(cols,start=1)})
    ws.freeze_panes=ws.cell(DATA0, 4)   # keep banner+headers and first 3 id columns
    last=DATA0+nrows-1
    ws.auto_filter.ref=f"A{HDR_S}:{L(ncols)}{last}"
    # column comments (glossary tooltips on a few key headers)
    tips={"MOQ":"Minimum Order Quantity — smallest run the factory will produce.",
          "Sample Lead":"Working days from tech pack to first sample.",
          "Prod. Lead":"Working days from PO to shipped bulk.",
          "Evaluation":"Your status for this supplier — pick from the dropdown.",
          "Next Step":"The single next action to move this supplier forward."}
    for name,tip in tips.items():
        cl=col_letter_by_name(cols,name)
        if cl: ws[f"{cl}{HDR_S}"].comment=Comment(tip,"BORN Studio")
    # data validation — Category
    if category_at:
        dv=DataValidation(type="list",
            formula1='"Knit,Woven,Activewear,Denim,Outerwear,Accessories,Trims & Notions"',
            allow_blank=True, showErrorMessage=True)
        dv.prompt="Pick a category"; dv.promptTitle="Category"
        ws.add_data_validation(dv); dv.add(f"{category_at}{DATA0}:{category_at}{last}")
    # data validation + conditional formatting — Evaluation
    if eval_at:
        dv=DataValidation(type="list", formula1='"Preferred,Shortlist,Hold,Pass"',
            allow_blank=True, showErrorMessage=True)
        dv.prompt="Set the supplier status"; dv.promptTitle="Evaluation"
        ws.add_data_validation(dv); dv.add(f"{eval_at}{DATA0}:{eval_at}{last}")
        rng=f"{eval_at}{DATA0}:{eval_at}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal",formula=['"Preferred"'],
            fill=fill(RED), font=Font(name=SANS,size=10,bold=True,color=WHITE)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal",formula=['"Shortlist"'],
            fill=fill(PAPER3), font=Font(name=SANS,size=10,bold=True,color=INK)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal",formula=['"Hold"'],
            fill=fill("FFF6E6C7"), font=Font(name=SANS,size=10,color=G2)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal",formula=['"Pass"'],
            fill=fill(PAPER2), font=Font(name=SANS,size=10,italic=True,color=G4)))
    build_guide(guide, banner_title, kicker, groups, glossary or [], has_eval=bool(eval_at))
    # print: landscape, fit all columns to one page wide
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_margins.left=ws.page_margins.right=ws.page_margins.top=ws.page_margins.bottom=0.3
    wb.save(path)
    return path

# ---------------------------------------------------------------- Start here sheet
def build_guide(ws, title, kicker, groups, glossary, has_eval):
    ws.sheet_view.showGridLines=False
    W=8
    set_widths(ws, {1:3,2:22,3:30,4:22,5:22,6:20,7:20,8:6})
    banner(ws, W, title, kicker)
    r=6
    ws.row_dimensions[r].height=10
    def heading(row, text):
        merge_center(ws,row,2,row,7,text,Font(name=SERIF,size=15,bold=True,color=INK),PAPER)
        ws.row_dimensions[row].height=26
    def body(row, text, font=None):
        merge_center(ws,row,2,row,7,text,font or Font(name=SANS,size=10,color=G2),PAPER,
                     Alignment(horizontal="left",vertical="center",wrap_text=True))
    r=7; heading(r,"How to use this tracker")
    steps=["1 — Each row is one factory. Start on the first data row; the example row shows the format.",
           "2 — Fill left to right: who they are, how they produce, then their costs per style.",
           "3 — Costs are in USD. Sample cost = one prototype; garment/production cost = per unit at MOQ.",
           "4 — Use the column filters (▾ on the header row) to compare, sort and shortlist.",
           ("5 — Set each supplier's status in Evaluation — Preferred turns Rojo." if has_eval
            else "5 — Keep one clear Next Step per supplier so nothing stalls.")]
    for i,s in enumerate(steps):
        rr=r+1+i; ws.row_dimensions[rr].height=22; body(rr,s)
    r=r+1+len(steps)+1
    heading(r,"The key")
    r+=1
    keys=[(INK,"Section header","Groups of columns — who / production / cost / status."),
          (PAPER3,"Column header","The field to fill. Hover for a tip on the key ones."),
          (PAPER2,"Your rows","Everything below the header is yours to complete."),
          (RED,"Preferred","In Evaluation, your chosen factories light up in Rojo Valentino.")]
    for kc,label,desc in keys:
        ws.row_dimensions[r].height=22
        ws.cell(r,2).fill=fill(kc); ws.cell(r,2).border=box(l=THINK,r=THINK,t=THINK,b=THINK)
        ws.cell(r,3).value=label; ws.cell(r,3).font=Font(name=SANS,size=10,bold=True,color=INK)
        ws.cell(r,3).alignment=Alignment(horizontal="left",vertical="center",indent=1)
        merge_center(ws,r,4,r,7,desc,Font(name=SANS,size=10,color=G2),PAPER)
        r+=1
    r+=1
    if glossary:
        heading(r,"Field glossary"); r+=1
        # header
        ws.cell(r,2).value="Field"; ws.cell(r,4).value="What to enter"
        for cc in (2,4):
            ws.cell(r,cc).font=Font(name=MONO,size=8,bold=True,color=G3)
        merge_center(ws,r,4,r,7,"What to enter",Font(name=MONO,size=8,bold=True,color=G3),PAPER)
        r+=1
        for field,desc in glossary:
            ws.row_dimensions[r].height=20
            ws.cell(r,2).value=field; ws.cell(r,2).font=Font(name=SANS,size=10,bold=True,color=INK)
            ws.cell(r,2).alignment=Alignment(horizontal="left",vertical="center")
            ws.cell(r,2).border=box(b=THIN)
            merge_center(ws,r,4,r,7,desc,Font(name=SANS,size=10,color=G2),PAPER)
            for c in range(4,8): ws.cell(r,c).border=box(b=THIN)
            ws.cell(r,3).border=box(b=THIN)
            r+=1
    r+=1
    merge_center(ws,r,2,r,7,"Born Studio · Full-Service Apparel Development · v1.0",
                 Font(name=MONO,size=8,color=G4),PAPER)
    # paint remaining background paper for a clean canvas
    for rr in range(5, r+2):
        for c in range(1,W+1):
            if ws.cell(rr,c).fill.patternType is None: ws.cell(rr,c).fill=fill(PAPER)

# =========================================================== SOURCING
styles_s=["Shams Wide-Leg Leggings","Majara Regular Leggings","Qamar Long line top",
          "Suha Longline Integrated Tank-Bra","Zohra Heroine Jacket","Noor Dress"]
cols_s=[("Country",14,None,"text"),("Category",16,None,"cat"),("Factory Name",30,None,"text"),
        ("Website",26,None,"link"),("Email",24,None,"link"),("POC",16,None,"text"),
        ("MOQ",10,'#,##0',"num"),("Sample Lead",12,'0" d"',"num"),("Prod. Lead",12,'0" d"',"num")]
cols_s+=[(s,15,'$#,##0',"cost") for s in styles_s]
cols_s+=[(s,15,'$#,##0',"cost") for s in styles_s]
cols_s+=[("Certifications",22,None,"text"),("Vendors Folder",18,None,"link"),
         ("Notes",42,None,"text"),("Evaluation",14,None,"eval")]
groups_s=[(1,3,"FACTORY"),(4,6,"CONTACT"),(7,9,"PRODUCTION"),
          (10,15,"SAMPLE COST · USD"),(16,21,"GARMENT COST · USD"),(22,25,"REFERENCE & STATUS")]
ex_s={"Country":"Portugal","Category":"Activewear","Factory Name":"Atelier Norte",
      "Website":"ateliernorte.pt","Email":"hello@ateliernorte.pt","POC":"Marta Sousa",
      "MOQ":300,"Sample Lead":15,"Prod. Lead":45,
      "Shams Wide-Leg Leggings":65,"Majara Regular Leggings":55,"Qamar Long line top":45,
      "Suha Longline Integrated Tank-Bra":60,"Zohra Heroine Jacket":120,"Noor Dress":85,
      "Certifications":"GOTS · OEKO-TEX","Vendors Folder":"▸ Drive link",
      "Notes":"Strong on technical knits; English-speaking POC.","Evaluation":"Preferred"}
# garment-cost example values (second block, same style names -> only first block set above;
# set garment block via positional patch after build not possible in dict; add distinct keys)
gloss_s=[("Country / Category","Where they are and what they specialise in."),
         ("Factory Name","Legal or trade name you'll reference everywhere."),
         ("Website / Email / POC","How to reach them and your point of contact."),
         ("MOQ","Minimum Order Quantity per style/colour."),
         ("Sample / Prod. Lead","Turnaround in working days."),
         ("Sample cost","Price of one prototype for each style."),
         ("Garment cost","Per-unit production price at MOQ for each style."),
         ("Certifications","GOTS, OEKO-TEX, BSCI, etc."),
         ("Vendors Folder","Link to their profile, quotes and docs."),
         ("Evaluation","Preferred · Shortlist · Hold · Pass.")]

# because the two cost blocks share style names, fill the garment example cells by index
def build_sourcing():
    p=f"{OUTDIR}/BORN_Sourcing_Tracker.xlsx"
    build_tracker(p,"Suppliers","SOURCING TRACKER","Supplier sourcing · v1.0",
                  groups_s, cols_s, ex_s, category_at="B", eval_at="Y",
                  glossary=gloss_s, nrows=40)
    # patch garment-cost example (columns 16..21) on the example row (row 7)
    wb=openpyxl.load_workbook(p); ws=wb["Suppliers"]
    gvals=[18,15,12,16,34,24]
    for i,v in enumerate(gvals): ws.cell(7,16+i).value=v
    wb.save(p); return p

# =========================================================== DEVELOPMENT
styles_d=["Mid range WOMEN","Mid Range MEN","Premium MEN","Premium silicone","Basic MEN",
          "Hoodie UNISEX","Zip pullover","Fitted Hat","6 panel Hat"]
cols_d=[("Factory Name",26,None,"text"),("Email",24,None,"link"),("POC",16,None,"text"),
        ("MOQ",10,'#,##0',"num"),("Sample Lead",12,'0" d"',"num"),("Prod. Lead",12,'0" d"',"num")]
cols_d+=[(s,14,'$#,##0',"cost") for s in styles_d]
cols_d+=[(s,14,'$#,##0',"cost") for s in styles_d]
cols_d+=[("Payment Terms",20,None,"text"),("Certifications",22,None,"text"),
         ("Invoice",14,None,"link"),("Samples Cost",14,'$#,##0',"cost"),("Next Step",24,None,"text")]
groups_d=[(1,3,"SUPPLIER"),(4,6,"PRODUCTION"),(7,15,"SAMPLE COST · USD"),
          (16,24,"PRODUCTION COST · USD"),(25,29,"TERMS & STATUS")]
ex_d={"Factory Name":"Atelier Norte","Email":"hello@ateliernorte.pt","POC":"Marta Sousa",
      "MOQ":300,"Sample Lead":15,"Prod. Lead":45,
      "Payment Terms":"30% deposit · Net 30","Certifications":"OEKO-TEX",
      "Invoice":"▸ link","Samples Cost":240,"Next Step":"Send tech packs"}
gloss_d=[("Factory Name","Legal or trade name."),
         ("Email / POC","Contact and your point of contact."),
         ("MOQ","Minimum Order Quantity per style."),
         ("Sample / Prod. Lead","Turnaround in working days."),
         ("Sample cost","Prototype price per style."),
         ("Production cost","Per-unit price at MOQ per style."),
         ("Payment Terms","Deposit % and net terms."),
         ("Certifications","OEKO-TEX, GRS, etc."),
         ("Invoice","Link to the latest invoice."),
         ("Samples Cost","Total spent on samples so far."),
         ("Next Step","The one action to move this factory forward.")]
def build_development():
    p=f"{OUTDIR}/BORN_Development_Tracker.xlsx"
    build_tracker(p,"Development","DEVELOPMENT TRACKER","Product development · v1.0",
                  groups_d, cols_d, ex_d, category_at=None, eval_at=None,
                  glossary=gloss_d, nrows=40)
    wb=openpyxl.load_workbook(p); ws=wb["Development"]
    svals=[45,40,70,95,32,38,42,18,22]; pvals=[14,12,22,30,9,12,13,6,7]
    for i,v in enumerate(svals): ws.cell(7,7+i).value=v
    for i,v in enumerate(pvals): ws.cell(7,16+i).value=v
    wb.save(p); return p

print(build_sourcing()); print(build_development())
